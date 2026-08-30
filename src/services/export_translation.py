"""Ручной экспорт поля перевода (UC-003, FT-004, FT-044, FT-046). Без виджетов."""

from __future__ import annotations

from pathlib import Path

from docx import Document

from domain.errors import ExportWriteError
from domain.models import ExportTranslationCommand

_ALLOWED_FORMATS = frozenset({"txt", "docx"})


def _write_txt(path: Path, text: str) -> None:
    path.write_text(text, encoding="utf-8", newline="\n")


def _write_docx(path: Path, text: str) -> None:
    document = Document()
    lines = text.split("\n")
    first = lines[0] if lines else ""
    if document.paragraphs:
        document.paragraphs[0].text = first
    else:
        document.add_paragraph(first)
    for line in lines[1:]:
        document.add_paragraph(line)
    document.save(str(path))


class ExportTranslation:
    """Пишет выбранный Пользователем файл. Workspace не очищает."""

    def run(self, command: ExportTranslationCommand) -> None:
        if len(command.translation_text) == 0:
            raise ExportWriteError("empty translation")
        if command.export_format not in _ALLOWED_FORMATS:
            raise ExportWriteError("unsupported format")
        path = Path(command.path)
        suffix = path.suffix.lower()
        if command.export_format == "txt" and suffix not in (".txt", ""):
            raise ExportWriteError("unsupported format")
        if command.export_format == "docx" and suffix not in (".docx", ""):
            raise ExportWriteError("unsupported format")
        if suffix == "":
            path = path.with_suffix(f".{command.export_format}")
        try:
            if command.export_format == "txt":
                _write_txt(path, command.translation_text)
            else:
                _write_docx(path, command.translation_text)
        except ExportWriteError:
            raise
        except Exception as exc:
            raise ExportWriteError(str(exc)) from exc
