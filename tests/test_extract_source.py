"""S-11/S-12: ExtractSource / LoadSource без GUI (FT-003, FT-039, FT-042)."""

from __future__ import annotations

import time
from collections.abc import Callable
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from domain.errors import DocumentParseError
from domain.models import LoadSourceCommand
from services.extract_source import ExtractSource
from ui.bridge import LoadSourceBridge
from use_cases.load_source import LoadSource


class FakeHost:
    def call_on_ui(self, func: Callable[[], None]) -> None:
        func()

    def winfo_exists(self) -> bool:
        return True


def _wait_until(predicate: Callable[[], bool]) -> None:
    deadline = time.monotonic() + 2.0
    while time.monotonic() < deadline:
        if predicate():
            return
        time.sleep(0.02)
    raise AssertionError("load source bridge did not finish")


def _write_docx(path: Path, text: str) -> None:
    document = Document()
    lines = text.split("\n")
    first = lines[0] if lines else ""
    if document.paragraphs:
        document.paragraphs[0].text = first
    else:
        document.add_paragraph(first)
    for line in lines[1:]:
        document.add_paragraph(line)
    document.save(str(path))


def _pdf_with_text(line: str) -> bytes:
    escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    bodies = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(bodies, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode("ascii")
        out += body
        out += b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(bodies) + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer << /Size {len(bodies) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


def test_should_return_utf8_text_when_txt_file_is_readable(tmp_path: Path) -> None:
    """FT-003: `.txt` UTF-8 попадает в исходник."""
    path = tmp_path / "in.txt"
    path.write_text("Hello\nвторая строка", encoding="utf-8")
    assert ExtractSource().run(str(path), "txt") == "Hello\nвторая строка"


def test_should_return_text_when_file_has_utf8_bom(tmp_path: Path) -> None:
    """FT-003: UTF-8 с BOM (`utf-8-sig`)."""
    path = tmp_path / "bom.txt"
    path.write_bytes(b"\xef\xbb\xbfBOM text")
    assert ExtractSource().run(str(path), "txt") == "BOM text"


def test_should_fallback_to_cp1251_when_utf8_fails(tmp_path: Path) -> None:
    """FT-003: после неудачного utf-8 — cp1251."""
    path = tmp_path / "win.txt"
    path.write_bytes("Привет".encode("cp1251"))
    assert ExtractSource().run(str(path), "txt") == "Привет"


def test_should_return_markdown_text_when_md_file_is_readable(tmp_path: Path) -> None:
    """FT-003: `.md` читается как текст."""
    path = tmp_path / "note.md"
    path.write_text("# Заголовок\n\nабзац", encoding="utf-8")
    assert ExtractSource().run(str(path), "md") == "# Заголовок\n\nабзац"


def test_should_raise_parse_error_when_file_is_empty(tmp_path: Path) -> None:
    """S-11: пустой `.txt` — не успех, поле не подменять."""
    path = tmp_path / "empty.txt"
    path.write_text("", encoding="utf-8")
    with pytest.raises(DocumentParseError):
        ExtractSource().run(str(path), "txt")


def test_should_raise_parse_error_when_file_is_only_whitespace(tmp_path: Path) -> None:
    """S-11: одни пробелы/переводы — текст не извлечён."""
    path = tmp_path / "blank.txt"
    path.write_text("  \n\n  ", encoding="utf-8")
    with pytest.raises(DocumentParseError):
        ExtractSource().run(str(path), "txt")


def test_should_return_paragraphs_when_docx_has_text(tmp_path: Path) -> None:
    """FT-003 / FT-042: `.docx` — абзацы в сырой текст."""
    path = tmp_path / "doc.docx"
    _write_docx(path, "Первый абзац\nВторой абзац")
    assert ExtractSource().run(str(path), "docx") == "Первый абзац\nВторой абзац"


def test_should_return_page_text_when_pdf_has_extractable_text(tmp_path: Path) -> None:
    """FT-003 / FT-042: `.pdf` с текстовым слоем."""
    path = tmp_path / "doc.pdf"
    path.write_bytes(_pdf_with_text("Hello PDF"))
    text = ExtractSource().run(str(path), "pdf")
    assert "Hello PDF" in text


def test_should_raise_parse_error_when_docx_is_empty(tmp_path: Path) -> None:
    """FT-039: пустой `.docx` — не успех."""
    path = tmp_path / "empty.docx"
    _write_docx(path, "")
    with pytest.raises(DocumentParseError):
        ExtractSource().run(str(path), "docx")


def test_should_raise_parse_error_when_pdf_has_no_text_layer(tmp_path: Path) -> None:
    """FT-039: PDF без текстового слоя (скан)."""
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(str(path))
    with pytest.raises(DocumentParseError):
        ExtractSource().run(str(path), "pdf")


def test_should_raise_parse_error_when_docx_is_corrupt(tmp_path: Path) -> None:
    """UC-002 E2: битый `.docx`."""
    path = tmp_path / "broken.docx"
    path.write_bytes(b"pk not a package")
    with pytest.raises(DocumentParseError):
        ExtractSource().run(str(path), "docx")


def test_should_raise_parse_error_when_pdf_is_corrupt(tmp_path: Path) -> None:
    """UC-002 E2: битый `.pdf`."""
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"%PDF-1.4 not a real file")
    with pytest.raises(DocumentParseError):
        ExtractSource().run(str(path), "pdf")


def test_should_return_text_when_load_source_opens_docx(tmp_path: Path) -> None:
    """LoadSource координирует ExtractSource для `.docx`."""
    path = tmp_path / "in.docx"
    _write_docx(path, "из документа")
    text = LoadSource().run(
        LoadSourceCommand(request_id=2, path=str(path), source_format="docx")
    )
    assert text == "из документа"


def test_should_raise_parse_error_when_file_cannot_be_read(tmp_path: Path) -> None:
    """UC-002 E1: ОС не отдаёт файл."""
    missing = tmp_path / "no-such.txt"
    with pytest.raises(DocumentParseError):
        ExtractSource().run(str(missing), "txt")


def test_should_keep_text_longer_than_100k_when_extracted(tmp_path: Path) -> None:
    """S-11: лимит 100k — на перевод, не на вместимость поля."""
    path = tmp_path / "huge.txt"
    text = "a" * 100_001
    path.write_text(text, encoding="utf-8")
    assert ExtractSource().run(str(path), "txt") == text


def test_should_return_text_when_load_source_use_case_runs(tmp_path: Path) -> None:
    """LoadSource координирует ExtractSource."""
    path = tmp_path / "in.txt"
    path.write_text("из use case", encoding="utf-8")
    text = LoadSource().run(
        LoadSourceCommand(request_id=1, path=str(path), source_format="txt")
    )
    assert text == "из use case"


def test_should_notify_success_when_load_bridge_reads(tmp_path: Path) -> None:
    """Клей: успешное чтение доходит до UI-колбэка."""
    host = FakeHost()
    successes: list[tuple[int, str]] = []
    errors: list[int] = []
    bridge = LoadSourceBridge(
        host=host,
        on_success=lambda rid, text: successes.append((rid, text)),
        on_error=lambda rid: errors.append(rid),
    )
    path = tmp_path / "bridge.txt"
    path.write_text("из моста", encoding="utf-8")
    bridge.start(path=str(path), source_format="txt")
    _wait_until(lambda: bool(successes) or bool(errors))
    assert successes == [(1, "из моста")]
    assert errors == []


def test_should_notify_error_when_load_bridge_parse_fails(tmp_path: Path) -> None:
    """Клей: пустой файл — on_error."""
    host = FakeHost()
    successes: list[tuple[int, str]] = []
    errors: list[int] = []
    bridge = LoadSourceBridge(
        host=host,
        on_success=lambda rid, text: successes.append((rid, text)),
        on_error=lambda rid: errors.append(rid),
    )
    path = tmp_path / "empty.txt"
    path.write_text("", encoding="utf-8")
    bridge.start(path=str(path), source_format="txt")
    _wait_until(lambda: bool(successes) or bool(errors))
    assert successes == []
    assert errors == [1]
