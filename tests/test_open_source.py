"""S-11/S-12: «Открыть файл» `.txt` / `.md` / `.docx` / `.pdf` (UC-002, FT-003, FT-039, FT-042, FT-047)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from ui.messages import (
    LABEL_OPEN_FILE,
    LABEL_ORIGINAL,
    LABEL_TRANSLATION,
    MSG_SOURCE_NOT_EXTRACTED,
    OPEN_FILETYPES,
    TITLE_SOURCE_NOT_EXTRACTED,
)
from ui_helpers import (
    find_button,
    find_textbox_for_label,
    pump_until,
    set_textbox_content,
    textbox_content,
)


def _fill_original(window, text: str) -> None:
    box = find_textbox_for_label(window, LABEL_ORIGINAL)
    assert box is not None
    set_textbox_content(box, text)
    box.event_generate("<KeyRelease>")
    window.update_idletasks()


def _fill_translation(window, text: str) -> None:
    box = find_textbox_for_label(window, LABEL_TRANSLATION)
    assert box is not None
    set_textbox_content(box, text)
    box.event_generate("<KeyRelease>")
    window.update_idletasks()


def _click_open(window) -> None:
    button = find_button(window, LABEL_OPEN_FILE)
    assert button is not None
    button.invoke()
    window.update_idletasks()


def _write_docx(path: Path, text: str) -> None:
    document = Document()
    lines = text.split("\n")
    first = lines[0] if lines else ""
    if document.paragraphs:
        document.paragraphs[0].text = first
    else:
        document.add_paragraph(first)
    for line in lines[1:]:
        document.add_paragraph(line)
    document.save(str(path))


def _pdf_with_text(line: str) -> bytes:
    escaped = line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
    bodies = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        ),
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(bodies, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode("ascii")
        out += body
        out += b"\nendobj\n"
    xref_at = len(out)
    out += f"xref\n0 {len(bodies) + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer << /Size {len(bodies) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_at}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


def _open_path(monkeypatch, window, path: Path) -> None:
    monkeypatch.setattr(
        "ui.layout.filedialog.askopenfilename",
        lambda **kwargs: str(path),
    )
    _click_open(window)


def test_should_put_txt_contents_in_original_when_user_opens_file(
    monkeypatch,
    window,
    tmp_path: Path,
):
    """FT-003: выбранный `.txt` появляется в «Оригинальный текст»."""
    target = tmp_path / "book.txt"
    target.write_text("Глава первая", encoding="utf-8")
    _open_path(monkeypatch, window, target)
    box = find_textbox_for_label(window, LABEL_ORIGINAL)
    assert box is not None
    pump_until(window, lambda: textbox_content(box) == "Глава первая")


def test_should_put_md_contents_in_original_when_user_opens_file(
    monkeypatch,
    window,
    tmp_path: Path,
):
    """FT-003: выбранный `.md` появляется в «Оригинальный текст»."""
    target = tmp_path / "note.md"
    target.write_text("# Note\n\nbody", encoding="utf-8")
    _open_path(monkeypatch, window, target)
    box = find_textbox_for_label(window, LABEL_ORIGINAL)
    assert box is not None
    pump_until(window, lambda: textbox_content(box) == "# Note\n\nbody")


def test_should_offer_txt_md_docx_pdf_when_open_dialog_opens(monkeypatch, window):
    """FT-042: диалог сразу показывает все четыре формата, без «все файлы»."""
    captured: dict[str, object] = {}

    def fake_dialog(**kwargs: object) -> str:
        captured.update(kwargs)
        return ""

    monkeypatch.setattr("ui.layout.filedialog.askopenfilename", fake_dialog)
    _click_open(window)
    assert captured.get("filetypes") == list(OPEN_FILETYPES)
    default_pattern = OPEN_FILETYPES[0][1]
    assert "*.txt" in default_pattern
    assert "*.md" in default_pattern
    assert "*.docx" in default_pattern
    assert "*.pdf" in default_pattern
    joined = " ".join(f"{name} {pattern}" for name, pattern in OPEN_FILETYPES)
    assert "*.*" not in joined and ".xlsx" not in joined


def test_should_keep_original_when_user_cancels_open_dialog(
    monkeypatch,
    window,
):
    """FT-047: отмена диалога не меняет левое поле."""
    monkeypatch.setattr(
        "ui.layout.filedialog.askopenfilename",
        lambda **kwargs: "",
    )
    _fill_original(window, "Был оригинал")
    _click_open(window)
    window.update_idletasks()
    box = find_textbox_for_label(window, LABEL_ORIGINAL)
    assert box is not None
    assert textbox_content(box) == "Был оригинал"


def test_should_keep_original_and_show_message_when_txt_is_empty(
    monkeypatch,
    window,
    tmp_path: Path,
):
    """S-11: пустой `.txt` — сообщение, поле не подменять молча."""
    target = tmp_path / "empty.txt"
    target.write_text("", encoding="utf-8")
    errors: list[tuple[str, str]] = []

    def fake_error(title: str, message: str, **kwargs: object) -> str:
        errors.append((title, message))
        return "ok"

    monkeypatch.setattr("ui.layout.messagebox.showerror", fake_error)
    _fill_original(window, "Старый оригинал")
    _open_path(monkeypatch, window, target)
    pump_until(window, lambda: bool(errors))
    box = find_textbox_for_label(window, LABEL_ORIGINAL)
    assert box is not None
    assert textbox_content(box) == "Старый оригинал"
    assert errors == [(TITLE_SOURCE_NOT_EXTRACTED, MSG_SOURCE_NOT_EXTRACTED)]


def test_should_clear_translation_when_open_file_succeeds(
    monkeypatch,
    window,
    tmp_path: Path,
):
    """A0099: успешная загрузка очищает поле перевода."""
    target = tmp_path / "src.txt"
    target.write_text("Новый исходник", encoding="utf-8")
    monkeypatch.setattr(
        "ui.layout.messagebox.askokcancel",
        lambda *_a, **_k: True,
    )
    _fill_original(window, "Старый оригинал")
    _fill_translation(window, "Старый перевод")
    _open_path(monkeypatch, window, target)
    original = find_textbox_for_label(window, LABEL_ORIGINAL)
    translation = find_textbox_for_label(window, LABEL_TRANSLATION)
    assert original is not None and translation is not None
    pump_until(window, lambda: textbox_content(original) == "Новый исходник")
    assert textbox_content(translation) == ""


def test_should_keep_translation_when_open_file_fails(
    monkeypatch,
    window,
    tmp_path: Path,
):
    """UC-002: сбой извлечения не очищает поле перевода."""
    target = tmp_path / "empty.txt"
    target.write_text("   ", encoding="utf-8")
    errors: list[tuple[str, str]] = []

    def fake_error(title: str, message: str, **kwargs: object) -> str:
        errors.append((title, message))
        return "ok"

    monkeypatch.setattr("ui.layout.messagebox.showerror", fake_error)
    monkeypatch.setattr(
        "ui.layout.messagebox.askokcancel",
        lambda *_a, **_k: True,
    )
    _fill_original(window, "Оригинал")
    _fill_translation(window, "Перевод")
    _open_path(monkeypatch, window, target)
    pump_until(window, lambda: bool(errors))
    original = find_textbox_for_label(window, LABEL_ORIGINAL)
    translation = find_textbox_for_label(window, LABEL_TRANSLATION)
    assert original is not None and translation is not None
    assert textbox_content(original) == "Оригинал"
    assert textbox_content(translation) == "Перевод"


def test_should_put_docx_contents_in_original_when_user_opens_file(
    monkeypatch,
    window,
    tmp_path: Path,
):
    """FT-003 / FT-042: выбранный `.docx` появляется в «Оригинальный текст»."""
    target = tmp_path / "book.docx"
    _write_docx(target, "Глава из Word")
    _open_path(monkeypatch, window, target)
    box = find_textbox_for_label(window, LABEL_ORIGINAL)
    assert box is not None
    pump_until(window, lambda: textbox_content(box) == "Глава из Word")


def test_should_put_pdf_contents_in_original_when_user_opens_file(
    monkeypatch,
    window,
    tmp_path: Path,
):
    """FT-003 / FT-042: выбранный `.pdf` с текстом появляется в поле."""
    target = tmp_path / "book.pdf"
    target.write_bytes(_pdf_with_text("Hello PDF"))
    _open_path(monkeypatch, window, target)
    box = find_textbox_for_label(window, LABEL_ORIGINAL)
    assert box is not None
    pump_until(window, lambda: "Hello PDF" in textbox_content(box))


def test_should_keep_original_and_show_message_when_pdf_has_no_text(
    monkeypatch,
    window,
    tmp_path: Path,
):
    """FT-039: PDF без текстового слоя — сообщение, поле не подменять."""
    target = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(str(target))
    errors: list[tuple[str, str]] = []

    def fake_error(title: str, message: str, **kwargs: object) -> str:
        errors.append((title, message))
        return "ok"

    monkeypatch.setattr("ui.layout.messagebox.showerror", fake_error)
    _fill_original(window, "Старый оригинал")
    _open_path(monkeypatch, window, target)
    pump_until(window, lambda: bool(errors))
    box = find_textbox_for_label(window, LABEL_ORIGINAL)
    assert box is not None
    assert textbox_content(box) == "Старый оригинал"
    assert errors == [(TITLE_SOURCE_NOT_EXTRACTED, MSG_SOURCE_NOT_EXTRACTED)]


def test_should_not_open_other_format_when_path_is_not_allowed(
    monkeypatch,
    window,
    tmp_path: Path,
):
    """FT-042: формат вне списка не открывать."""
    target = tmp_path / "sheet.xlsx"
    target.write_bytes(b"not a spreadsheet")
    errors: list[tuple[str, str]] = []

    def fake_error(title: str, message: str, **kwargs: object) -> str:
        errors.append((title, message))
        return "ok"

    monkeypatch.setattr("ui.layout.messagebox.showerror", fake_error)
    _fill_original(window, "Был оригинал")
    _open_path(monkeypatch, window, target)
    window.update_idletasks()
    box = find_textbox_for_label(window, LABEL_ORIGINAL)
    assert box is not None
    assert textbox_content(box) == "Был оригинал"
    assert errors == [(TITLE_SOURCE_NOT_EXTRACTED, MSG_SOURCE_NOT_EXTRACTED)]
